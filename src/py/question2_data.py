"""Validated annual input and time mapping for Question 2."""

from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime, timedelta
from pathlib import Path

import numpy as np
from openpyxl import load_workbook

from question1 import Question1Data, load_question1_data


SLOTS_PER_DAY = 144
SLOT_MINUTES = 10


@dataclass(frozen=True)
class YearData:
    dates: tuple[date, ...]
    minute_of_day: np.ndarray
    price_yuan_per_kwh: np.ndarray
    load_kw: np.ndarray
    pv_kw: np.ndarray

    @property
    def load_kwh(self) -> np.ndarray:
        return self.load_kw / 6.0

    @property
    def pv_kwh(self) -> np.ndarray:
        return self.pv_kw / 6.0


@dataclass(frozen=True)
class ColdStartForecast:
    load_kw: np.ndarray
    pv_kw: np.ndarray


def _as_date(value: object) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return datetime.fromisoformat(str(value).strip()).date()


def _validate_matrix(values: np.ndarray, name: str) -> None:
    if values.ndim != 2 or values.shape[1] != SLOTS_PER_DAY:
        raise ValueError(f"{name} must have shape (days, {SLOTS_PER_DAY})")
    if not np.all(np.isfinite(values)) or np.any(values < 0.0):
        raise ValueError(f"{name} contains a missing, non-finite or negative value")


def _read_sheet(path: Path, sheet_name: str) -> tuple[tuple[date, ...], tuple[str, ...], np.ndarray]:
    workbook = load_workbook(path, data_only=True, read_only=True)
    sheet = workbook[sheet_name]
    rows_iterator = sheet.iter_rows(values_only=True)
    header_row = next(rows_iterator)
    headers = tuple(str(value).strip() for value in header_row[1:])
    if len(headers) != SLOTS_PER_DAY or len(set(headers)) != SLOTS_PER_DAY:
        raise ValueError(f"{sheet_name} must contain 144 unique time columns")
    dates: list[date] = []
    rows: list[list[float]] = []
    for row in rows_iterator:
        raw_date = row[0]
        if raw_date is None:
            continue
        dates.append(_as_date(raw_date))
        if len(row[1:]) != SLOTS_PER_DAY or any(value is None for value in row[1:]):
            raise ValueError(f"{sheet_name} has an incomplete row for {raw_date}")
        rows.append([float(value) for value in row[1:]])
    values = np.asarray(rows, dtype=float)
    _validate_matrix(values, sheet_name)
    return tuple(dates), headers, values


def load_year_data(attachment1: Path, attachment2: Path) -> YearData:
    """Read 2025 actual curves and rotate the final source column to midnight."""

    q1 = load_question1_data(attachment1)
    load_dates, load_headers, source_load = _read_sheet(attachment2, "小区负载")
    pv_dates, pv_headers, source_pv = _read_sheet(attachment2, "光伏发电实际功率")
    if load_dates != pv_dates or load_headers != pv_headers:
        raise ValueError("Load and PV sheets do not use identical dates and time columns")
    expected = tuple(date(2025, 1, 1) + timedelta(days=index) for index in range(365))
    if load_dates != expected or len(set(load_dates)) != 365:
        raise ValueError("Attachment 2 must contain the complete ordered 2025 calendar")
    # Source order is 00:10,...,23:50,00:00+1. Internal order is 00:00,...,23:50.
    load_kw = np.concatenate((source_load[:, -1:], source_load[:, :-1]), axis=1)
    pv_kw = np.concatenate((source_pv[:, -1:], source_pv[:, :-1]), axis=1)
    _validate_matrix(load_kw, "load_kw")
    _validate_matrix(pv_kw, "pv_kw")
    return YearData(
        dates=load_dates,
        minute_of_day=np.arange(0, 1440, SLOT_MINUTES, dtype=int),
        price_yuan_per_kwh=np.asarray(q1.price_yuan_per_kwh, dtype=float),
        load_kw=load_kw,
        pv_kw=pv_kw,
    )


def load_cold_start_forecast(path: Path) -> ColdStartForecast:
    data: Question1Data = load_question1_data(path)
    return ColdStartForecast(data.load_kw.copy(), data.pv_kw.copy())


def template_column_minutes() -> np.ndarray:
    """Return internal slot starts in the official 00:10,...,00:00+1 order."""

    return np.concatenate((np.arange(10, 1440, 10, dtype=int), np.asarray([0], dtype=int)))


def internal_to_template(values: np.ndarray) -> np.ndarray:
    values = np.asarray(values, dtype=float)
    if values.shape != (SLOTS_PER_DAY,):
        raise ValueError("Expected one 144-slot day")
    return np.concatenate((values[1:], values[:1]))
