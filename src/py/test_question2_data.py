"""Tests for Question 2 annual data and time mapping."""

from __future__ import annotations

import unittest
from pathlib import Path

import numpy as np
from openpyxl import load_workbook

from question2_data import (
    internal_to_template,
    load_actual_year_data,
    load_year_data,
    template_column_minutes,
)


ROOT = Path(__file__).resolve().parents[2]
ATTACHMENT_1 = ROOT / "problems/C题/附件/附件1.xlsx"
ATTACHMENT_2 = ROOT / "problems/C题/附件/附件2.xlsx"


class Question2DataTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.data = load_year_data(ATTACHMENT_1, ATTACHMENT_2)

    def test_year_has_complete_calendar_and_slots(self) -> None:
        self.assertEqual(str(self.data.dates[0]), "2025-01-01")
        self.assertEqual(str(self.data.dates[-1]), "2025-12-31")
        self.assertEqual(self.data.load_kw.shape, (365, 144))
        self.assertEqual(self.data.pv_kw.shape, (365, 144))
        np.testing.assert_array_equal(self.data.minute_of_day, np.arange(0, 1440, 10))

    def test_last_source_column_rotates_to_midnight(self) -> None:
        source = load_workbook(ATTACHMENT_2, read_only=True, data_only=True)["小区负载"]
        self.assertEqual(self.data.load_kw[0, 0], source.cell(2, 145).value)
        self.assertEqual(self.data.load_kw[0, 1], source.cell(2, 2).value)

    def test_actual_year_loader_matches_question2_curves(self) -> None:
        actual = load_actual_year_data(ATTACHMENT_2)
        self.assertEqual(actual.dates, self.data.dates)
        np.testing.assert_array_equal(actual.minute_of_day, self.data.minute_of_day)
        np.testing.assert_allclose(actual.load_kw, self.data.load_kw)
        np.testing.assert_allclose(actual.pv_kw, self.data.pv_kw)

    def test_template_round_trip(self) -> None:
        internal = np.arange(144, dtype=float)
        template = internal_to_template(internal)
        self.assertEqual(template[0], 1.0)
        self.assertEqual(template[-1], 0.0)
        np.testing.assert_array_equal(template_column_minutes()[:2], np.array([10, 20]))
        self.assertEqual(template_column_minutes()[-1], 0)


if __name__ == "__main__":
    unittest.main()
