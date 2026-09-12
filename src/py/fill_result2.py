"""把问题二结果回填到官方模板 result2.xlsx。

用法（在仓库根目录执行）：
    python src/py/fill_result2.py                 # 变体 A：严格按计划执行（默认）
    python src/py/fill_result2.py --variant B     # 变体 B：购电量照付不议 + 储能日内再调度
    python src/py/fill_result2.py --limit 20      # 只填前 20 天，用于检查格式

本文件只保留旧模型复现实验，默认输出 ``src/outputs/result2_legacy.xlsx``，不会覆盖正式
工作簿。正式结果请运行 ``question2.py --write-results``。

三个工作表按题目要求填写
------------------------
1) 计划购电量：334 天（2025-02-01 ~ 12-31）× 144 个 10 分钟时段的计划购电量（kWh），
   并填写每日购电量、购电费汇总
2) 充放电量：每天 6 个四小时时段（0:00-4:00 ... 20:00-24:00）的充放电量，
   以及 0:00 与 24:00 的储电量。日期只写在当天首行，0:00 储电量写在当天第 1 行、
   24:00 储电量写在当天第 2 行
3) 紧急购电量：日期写在当天首行；同一段连续的紧急购电时段合并为一行（题目表 4 的写法）

口径
----
- 时间戳 = 时段起点，0:00+1 ≡ 0:00；列标签 "0:10-0:20" 对应时段 j=1，"0:00+1-0:10+1" 对应 j=0
- 计划购电量按题目"照付不议"结算，实际净负荷超出"计划购电量 + 计划放电 - 计划充电"的部分
  按 5 倍价紧急购电
- 输出保留 4 位小数（按仓库"模型口径"工作表的规定），比较方案时用未四舍五入的内部值
"""
from __future__ import annotations

import argparse
import copy
import datetime as dt
from pathlib import Path

import numpy as np
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter

import q2_model as q
import q2_stochastic as qs
from workbook_style import normalize_populated_fonts

BLOCKS = [(0, 240), (240, 480), (480, 720), (720, 960), (960, 1200), (1200, 1440)]
DEC = 4


# --------------------------------------------------------------------------- 时间工具
def fmt_clock(minute: int) -> str:
    """分钟 -> 模板风格的时刻文本；1440 写作 0:00+1。"""
    day, m = divmod(int(minute), 1440)
    h, mm = divmod(m, 60)
    return f"{h}:{mm:02d}" + (f"+{day}" if day else "")


def label_to_period(label: str) -> int:
    """列标签（如 "0:10-0:20"）-> 时段序号 j（0..143）。"""
    start = str(label).strip().split("-")[0]
    offset = 1440 if start.endswith("+1") else 0
    if offset:
        start = start[:-2]
    h, m = start.split(":")
    return ((offset + int(h) * 60 + int(m)) % 1440) // 10


def block_label(start_min: int) -> str:
    return f"{fmt_clock(start_min)}-{fmt_clock(start_min + 240)}"


def cell_date(v) -> dt.date:
    """单元格日期：兼容 datetime 对象与 Excel 序列号两种写法。"""
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    return dt.date(1899, 12, 30) + dt.timedelta(days=int(v))


# --------------------------------------------------------------------------- 逐日结果
def compute_days(att: q.Attachment, variant: str, limit: int | None = None) -> list[dict]:
    days = []
    e = q.E0_KWH
    window = att.window[:limit] if limit else att.window
    for i in window:
        L, V = att.load_kwh[i], att.pv_kwh[i]
        net_act = L - V
        fl, fv = q.forecast_day(att, i)
        plan = q.plan_lp(att.price, fl, fv, e, "cycle")
        if variant == "B":
            rt = qs.realtime_dispatch(att, plan["g"], L, V, e, float(plan["e"][-1]))
            charge, discharge = rt["c"], rt["d"]
            soc_end = float(rt["e"][-1])
            emg_check = rt["emg"]
        else:
            charge, discharge = plan["c"], plan["d"]
            soc_end = float(plan["e"][-1])
            emg_check = None
        committed = plan["g"] + discharge - charge
        shortfall = np.maximum(0.0, net_act - committed)
        if emg_check is not None:      # 自检：变体 B 下两种算法必须一致
            assert np.abs(shortfall - emg_check).max() < 1e-6, "紧急购电量两种算法不一致"
        days.append({
            "date": att.dates[i], "index": i,
            "g": plan["g"], "charge": charge, "discharge": discharge,
            "soc_start": e, "soc_end": soc_end,
            "shortfall": shortfall,
            "plan_cost": float(np.dot(att.price, plan["g"])),
            "emg_cost": float(5 * np.dot(att.price, shortfall)),
        })
        e = soc_end
    return days


def merge_intervals(shortfall: np.ndarray, tol: float = 1e-6):
    """把连续的紧急购电时段合并成 (起始时段, 结束时段, 电量) 列表。"""
    out, t = [], 0
    while t < q.N:
        if shortfall[t] > tol:
            t0 = t
            while t < q.N and shortfall[t] > tol:
                t += 1
            out.append((t0, t, float(shortfall[t0:t].sum())))
        else:
            t += 1
    return out


# --------------------------------------------------------------------------- 写入
def style_from(ws, src_row: int, dst_row: int, cols) -> None:
    """把某一行的单元格格式复制到目标行（新增行保持与模板一致的格式）。"""
    for c in cols:
        s, d = ws.cell(src_row, c), ws.cell(dst_row, c)
        d.number_format = s.number_format
        d.font = copy.copy(s.font)
        d.border = copy.copy(s.border)
        d.alignment = copy.copy(s.alignment)


def fill(template: Path, output: Path, days: list[dict], variant: str) -> None:
    wb = load_workbook(template)
    for name in ("计划购电量", "充放电量", "紧急购电量"):
        if name not in wb.sheetnames:
            raise ValueError(f"模板缺少工作表 {name}：{wb.sheetnames}")

    # ---------------- 1) 计划购电量
    ws = wb["计划购电量"]
    header = {c: ws.cell(1, c).value for c in range(2, 2 + q.N)}
    col_period = {}
    for c, label in header.items():
        if label is None:
            raise ValueError(f"模板表头第 1 行第 {c} 列缺少列标签")
        col_period[c] = label_to_period(label)
    if sorted(col_period.values()) != list(range(q.N)):
        raise ValueError("模板列标签与时段序号不是一一对应")
    serial_to_row = {cell_date(ws.cell(r, 1).value): r for r in range(2, ws.max_row + 1)
                     if ws.cell(r, 1).value is not None}
    for d in days:
        if d["date"] not in serial_to_row:
            raise ValueError(f"模板计划购电量工作表缺少日期 {d['date']}")
        row = serial_to_row[d["date"]]
        for c, j in col_period.items():
            cell = ws.cell(row, c)
            cell.value = round(float(d["g"][j]), DEC)
            cell.number_format = "0.0000"
        # 模板最后两列分别为全天购电量与全天购电费。旧版填表器只写入
        # 144 个时段，导致问题四正式结果的日汇总列为空。
        ws.cell(row, 146).value = round(float(d["g"].sum()), DEC)
        ws.cell(row, 147).value = round(float(d["plan_cost"]), DEC)
        ws.cell(row, 146).number_format = "0.0000"
        ws.cell(row, 147).number_format = "0.0000"

    # ---------------- 2) 充放电量
    ws = wb["充放电量"]
    block_labels = [ws.cell(r, 2).value for r in range(2, 2 + len(BLOCKS))]
    if any(v is None for v in block_labels):
        raise ValueError("模板充放电量工作表缺少四小时时段标签")
    r = 2
    for d in days:
        for k, (a, b) in enumerate(BLOCKS):
            row = r + k
            if row > ws.max_row:
                style_from(ws, 2, row, range(1, 7))
            ws.cell(row, 1).value = d["date"] if k == 0 else None
            ws.cell(row, 2).value = block_labels[k]
            ws.cell(row, 3).value = round(float(d["charge"][a // 10:b // 10].sum()), DEC)
            ws.cell(row, 4).value = round(float(d["discharge"][a // 10:b // 10].sum()), DEC)
            ws.cell(row, 3).number_format = "0.0000"
            ws.cell(row, 4).number_format = "0.0000"
        ws.cell(r, 5).value = "0:00"
        ws.cell(r, 6).value = round(float(d["soc_start"]), DEC)
        ws.cell(r + 1, 5).value = "24:00"
        ws.cell(r + 1, 6).value = round(float(d["soc_end"]), DEC)
        ws.cell(r, 6).number_format = "0.0000"
        ws.cell(r + 1, 6).number_format = "0.0000"
        r += len(BLOCKS)
    last_row = r - 1
    for row in range(r, ws.max_row + 1):       # 清掉模板里的占位行
        for c in range(1, 7):
            ws.cell(row, c).value = None

    # ---------------- 3) 紧急购电量
    ws = wb["紧急购电量"]
    r = 2
    for d in days:
        items = merge_intervals(d["shortfall"])
        row = r
        if row > ws.max_row:
            style_from(ws, 2, row, range(1, 4))
        ws.cell(row, 1).value = d["date"]
        if items:
            t0, t1, energy = items[0]
            ws.cell(row, 2).value = f"{fmt_clock(10 * t0)}-{fmt_clock(10 * t1)}"
            ws.cell(row, 3).value = round(energy, DEC)
            ws.cell(row, 3).number_format = "0.0000"
        for t0, t1, energy in items[1:]:
            row += 1
            if row > ws.max_row:
                style_from(ws, 3, row, range(1, 4))
            ws.cell(row, 1).value = None
            ws.cell(row, 2).value = f"{fmt_clock(10 * t0)}-{fmt_clock(10 * t1)}"
            ws.cell(row, 3).value = round(energy, DEC)
            ws.cell(row, 3).number_format = "0.0000"
        r = row + 1
    for row in range(r, ws.max_row + 1):       # 清掉模板里的占位行
        for c in range(1, 4):
            ws.cell(row, c).value = None

    output.parent.mkdir(parents=True, exist_ok=True)
    normalize_populated_fonts(wb)
    wb.save(output)


# --------------------------------------------------------------------------- 校验
def verify(output: Path, days: list[dict]) -> None:
    def parse_ref(ref: str):
        col = "".join(ch for ch in ref if ch.isalpha())
        row = int("".join(ch for ch in ref if ch.isdigit()))
        return col, row

    grid = q.read_workbook(output)
    g_sheet = grid["计划购电量"]
    a_rows = {parse_ref(k)[1] for k, v in g_sheet.items()
              if parse_ref(k)[0] == "A" and str(v).replace("/", "").isdigit()}
    total_g = sum(float(v) for k, v in g_sheet.items()
                  if 2 <= column_index_from_string(parse_ref(k)[0]) <= 145
                  and parse_ref(k)[1] >= 2 and v not in ("", None))
    total_daily_g = sum(float(v) for k, v in g_sheet.items()
                        if parse_ref(k)[0] == "EP" and parse_ref(k)[1] >= 2
                        and v not in ("", None))
    total_daily_cost = sum(float(v) for k, v in g_sheet.items()
                           if parse_ref(k)[0] == "EQ" and parse_ref(k)[1] >= 2
                           and v not in ("", None))
    print(f"校验：计划购电量工作表数据行数 = {len([r for r in a_rows if r >= 2])}（应为 334）")
    print(f"      计划购电量合计 = {total_g:,.1f} kWh（内部值 "
          f"{sum(d['g'].sum() for d in days):,.1f}）")
    print(f"      日汇总列合计 = {total_daily_g:,.1f} kWh / {total_daily_cost:,.1f} 元（内部值 "
          f"{sum(d['g'].sum() for d in days):,.1f} / "
          f"{sum(d['plan_cost'] for d in days):,.1f}）")
    if abs(total_daily_g - sum(d["g"].sum() for d in days)) > 0.1:
        raise ValueError("计划购电量工作表的全天购电量汇总不一致")
    if abs(total_daily_cost - sum(d["plan_cost"] for d in days)) > 0.1:
        raise ValueError("计划购电量工作表的全天购电费汇总不一致")

    e_sheet = grid["紧急购电量"]
    emg_cells = [(k, v) for k, v in e_sheet.items() if parse_ref(k)[0] == "C"
                 and parse_ref(k)[1] >= 2 and v not in ("", None)]
    print(f"      紧急购电行数 = {len(emg_cells)}，合计 = "
          f"{sum(float(v) for _, v in emg_cells):,.1f} kWh（内部值 "
          f"{sum(d['shortfall'].sum() for d in days):,.1f}）")

    c_sheet = grid["充放电量"]
    ch = sum(float(v) for k, v in c_sheet.items() if parse_ref(k)[0] == "C"
             and parse_ref(k)[1] >= 2 and v not in ("", None))
    dis = sum(float(v) for k, v in c_sheet.items() if parse_ref(k)[0] == "D"
              and parse_ref(k)[1] >= 2 and v not in ("", None))
    print(f"      充放电量工作表：充电合计 {ch:,.1f} / 放电合计 {dis:,.1f} kWh（内部值 "
          f"{sum(d['charge'].sum() for d in days):,.1f} / "
          f"{sum(d['discharge'].sum() for d in days):,.1f}）")

    print("      抽查：")
    d0 = days[0]
    print(f"        {d0['date']} 计划购电量前 4 段 = "
          f"{[round(float(v), 4) for v in d0['g'][:4]]}")
    for r in range(2, 8):
        print(f"        充放电量 r{r}: " + " | ".join(
            str(c_sheet.get(f"{c}{r}", "")) for c in "ABCDEF"))
    for r in range(2, 6):
        print(f"        紧急购电量 r{r}: " + " | ".join(
            str(e_sheet.get(f"{c}{r}", "")) for c in "ABC"))
    print(f"      总计 {len(days)} 天：计划购电费 {sum(d['plan_cost'] for d in days):,.1f} 元，"
          f"紧急购电费 {sum(d['emg_cost'] for d in days):,.1f} 元")


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--variant", choices=["A", "B"], default="A")
    ap.add_argument("--limit", type=int, default=None, help="只填前 N 天（检查格式用）")
    ap.add_argument("--data-root", type=Path, default=None)
    ap.add_argument("--template", type=Path, default=None)
    ap.add_argument("--output", type=Path, default=None)
    args = ap.parse_args()

    root = Path(args.data_root) if args.data_root else q.project_root()
    template = Path(args.template) if args.template else \
        root / "problems" / "C题" / "附件" / "附件5" / "result2.xlsx"
    output = Path(args.output) if args.output else root / "src" / "outputs" / "result2_legacy.xlsx"
    official = (root / "src" / "附件5" / "result2.xlsx").resolve()
    if output.resolve() == official:
        raise SystemExit("旧模型入口禁止覆盖正式 result2.xlsx；请运行 question2.py --write-results")

    att = q.Attachment(root)
    days = compute_days(att, args.variant, args.limit)
    fill(template, output, days, args.variant)
    print(f"已写入 {output}（口径：变体 {args.variant}，{len(days)} 天）")
    verify(output, days)


if __name__ == "__main__":
    main()
