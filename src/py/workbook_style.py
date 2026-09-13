"""统一正式结果工作簿中已填充单元格的字体。"""
from __future__ import annotations

import copy
import os
import tempfile
from pathlib import Path

from openpyxl import load_workbook


RESULT_FONT_NAME = "宋体"
RESULT_FONT_SIZE = 10.0


def normalize_populated_fonts(workbook) -> None:
    """统一已用区域字体，并清除会让 Excel 回退到 Calibri 的主题字体属性。"""

    for worksheet in workbook.worksheets:
        for row in worksheet.iter_rows():
            for cell in row:
                font = copy.copy(cell.font)
                font.name = RESULT_FONT_NAME
                font.sz = RESULT_FONT_SIZE
                font.scheme = None
                font.family = None
                font.charset = None
                cell.font = font


def normalize_file(path: Path) -> None:
    """原子化更新一个工作簿，避免中途失败损坏正式结果。"""

    workbook = load_workbook(path)
    normalize_populated_fonts(workbook)
    descriptor, temporary_name = tempfile.mkstemp(suffix=".xlsx", dir=path.parent)
    os.close(descriptor)
    temporary = Path(temporary_name)
    try:
        workbook.save(temporary)
        workbook.close()
        os.replace(temporary, path)
    finally:
        temporary.unlink(missing_ok=True)


def main() -> None:
    root = Path(__file__).resolve().parents[2]
    result_dir = root / "src" / "附件5"
    for name in ("result1.xlsx", "result2.xlsx", "result3.xlsx", "result4-2.xlsx", "result4-3.xlsx"):
        path = result_dir / name
        if not path.exists():
            raise FileNotFoundError(path)
        normalize_file(path)
        print(f"字体已统一：{path}")


if __name__ == "__main__":
    main()
