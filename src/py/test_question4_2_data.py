"""Tests for Question 4-2 variable-price data loading."""

from __future__ import annotations

import shutil
import tempfile
import unittest
from pathlib import Path

import numpy as np
from openpyxl import load_workbook

from question4_2_data import load_question42_data, summarize_prices


ROOT = Path(__file__).resolve().parents[2]
ATTACHMENT_2 = ROOT / "problems/C题/附件/附件2.xlsx"
ATTACHMENT_4 = ROOT / "problems/C题/附件/附件4.xlsx"


class Question42DataTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.data = load_question42_data(ATTACHMENT_2, ATTACHMENT_4)

    def test_price_dates_slots_and_rotation_match_source(self) -> None:
        self.assertEqual(str(self.data.dates[0]), "2025-01-01")
        self.assertEqual(str(self.data.dates[-1]), "2025-12-31")
        self.assertEqual(self.data.price_yuan_per_kwh.shape, (365, 144))
        np.testing.assert_array_equal(self.data.minute_of_day, np.arange(0, 1440, 10))

        source = load_workbook(ATTACHMENT_4, read_only=True, data_only=True).active
        self.assertEqual(self.data.price_yuan_per_kwh[0, 0], source.cell(2, 145).value)
        self.assertEqual(self.data.price_yuan_per_kwh[0, 1], source.cell(2, 2).value)

    def test_price_summary_matches_hand_checked_source_controls(self) -> None:
        summary = summarize_prices(self.data)
        self.assertAlmostEqual(summary["minimum_yuan_per_kwh"], 0.0076)
        self.assertAlmostEqual(summary["maximum_yuan_per_kwh"], 1.7936)
        self.assertAlmostEqual(summary["mean_yuan_per_kwh"], 0.7661976046423135)
        self.assertGreater(summary["lag1_curve_correlation"], 0.97)
        self.assertGreater(summary["lag7_curve_correlation"], 0.98)

    def test_price_dates_must_match_actual_curve_dates(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            changed = Path(directory) / "attachment4_bad_dates.xlsx"
            shutil.copy2(ATTACHMENT_4, changed)
            workbook = load_workbook(changed)
            sheet = workbook.active
            sheet.cell(3, 1).value = sheet.cell(2, 1).value
            workbook.save(changed)

            with self.assertRaisesRegex(ValueError, "Attachment 4 dates must match Attachment 2"):
                load_question42_data(ATTACHMENT_2, changed)


if __name__ == "__main__":
    unittest.main()
