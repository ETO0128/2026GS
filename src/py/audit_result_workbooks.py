"""Audit official result workbooks without modifying them.

Checks populated-cell font consistency and reports formulas and worksheet sizes.
Run after any result workbook is regenerated.
"""
from __future__ import annotations

from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

import q2_model as q2


EXPECTED_FONT = ("宋体", 10.0)
FILES = ("result1.xlsx", "result2.xlsx", "result3.xlsx", "result4-2.xlsx", "result4-3.xlsx")


def audit(path: Path) -> list[str]:
    wb = load_workbook(path, data_only=False, read_only=False)
    messages: list[str] = []
    for ws in wb.worksheets:
        styles: Counter[tuple[str | None, float | None]] = Counter()
        formula_count = 0
        populated = 0
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                populated += 1
                styles[(cell.font.name, float(cell.font.sz) if cell.font.sz else None)] += 1
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    formula_count += 1
        bad = sum(count for style, count in styles.items() if style != EXPECTED_FONT)
        messages.append(
            f"{path.name}/{ws.title}: {ws.max_row}x{ws.max_column}, "
            f"populated={populated}, formulas={formula_count}, off-style={bad}, styles={dict(styles)}"
        )
        if bad:
            raise AssertionError(f"{path.name}/{ws.title} contains {bad} populated cells outside {EXPECTED_FONT}")
        if ws.max_column == 147 and ws.max_row == 335:
            dates = [ws.cell(row, 1).value for row in range(2, 336)]
            if any(value is None for value in dates) or len(set(dates)) != 334:
                raise AssertionError(f"{path.name}/{ws.title} does not contain 334 unique dates")
            quantity_errors = []
            for row in range(2, 336):
                period_total = sum(float(ws.cell(row, col).value or 0.0) for col in range(2, 146))
                reported_total = float(ws.cell(row, 146).value)
                if abs(period_total - reported_total) > 0.02:
                    quantity_errors.append(row)
            if quantity_errors:
                raise AssertionError(
                    f"{path.name}/{ws.title} has inconsistent daily quantities in rows {quantity_errors[:5]}"
                )
            messages.append(f"{path.name}/{ws.title}: 334 dates and daily quantity totals verified")
    return messages


def main() -> None:
    result_dir = q2.project_root() / "src" / "附件5"
    for name in FILES:
        for line in audit(result_dir / name):
            print(line)
    print("All populated result cells use 宋体 10 pt.")


if __name__ == "__main__":
    main()
