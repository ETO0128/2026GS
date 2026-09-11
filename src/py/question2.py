"""Run the causal annual forecast--dispatch--settlement loop for Question 2."""

from __future__ import annotations

import argparse
import os
import tempfile
import time
from copy import copy
from dataclasses import dataclass, replace
from datetime import date
from pathlib import Path

import numpy as np
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from question2_data import ColdStartForecast, YearData, internal_to_template, load_cold_start_forecast, load_year_data
from question2_dispatch import (
    DayAheadPlan,
    DayExecution,
    RealizedDay,
    compress_emergency_events,
    plan_day_ahead,
    plan_two_stage_stochastic,
    simulate_fixed_plan,
)
from question2_forecast import ForecastConfig, ForecastResult, forecast_day, forecast_metrics
from question2_scenarios import build_joint_residual_scenarios


INITIAL_SOC_KWH = 6000.0
OFFICIAL_START = date(2025, 2, 1)
SPECIFIED_DATES = (date(2025, 3, 20), date(2025, 6, 21), date(2025, 9, 23), date(2025, 12, 21))


@dataclass(frozen=True)
class Question2Config:
    forecast: ForecastConfig = ForecastConfig()
    planner: str = "deterministic"
    reserve_kwh: float = 6000.0
    emergency_multiplier: float = 5.0
    scenario_count: int = 14
    scenario_decay: float = 0.95
    scenario_window_days: int = 90


@dataclass(frozen=True)
class DailyRecord:
    date: date
    forecast: ForecastResult
    plan: DayAheadPlan
    execution: DayExecution
    actual_load_kw: np.ndarray
    actual_pv_kw: np.ndarray
    expected_emergency_cost_yuan: float = 0.0


@dataclass(frozen=True)
class YearResult:
    days: tuple[DailyRecord, ...]
    metrics: dict[str, float]
    run_metadata: dict[str, str | float | int]

    @property
    def official_days(self) -> tuple[DailyRecord, ...]:
        return tuple(day for day in self.days if day.date >= OFFICIAL_START)


def calibrate_stochastic_parameters(
    data: YearData,
    cold_start: ColdStartForecast,
    base_config: Question2Config,
    candidate_counts: tuple[int, ...] = (7, 14, 28),
    candidate_decays: tuple[float, ...] = (0.90, 0.95, 0.98),
) -> tuple[Question2Config, dict[str, float]]:
    """Select scenario count and decay using January only."""

    january_days = sum(item.month == 1 for item in data.dates)
    january = YearData(
        dates=data.dates[:january_days],
        minute_of_day=data.minute_of_day.copy(),
        price_yuan_per_kwh=data.price_yuan_per_kwh.copy(),
        load_kw=data.load_kw[:january_days].copy(),
        pv_kw=data.pv_kw[:january_days].copy(),
    )
    scores: dict[str, float] = {}
    candidates: list[tuple[float, int, float, Question2Config]] = []
    for count in candidate_counts:
        for decay in candidate_decays:
            config = replace(
                base_config,
                planner="stochastic",
                scenario_count=count,
                scenario_decay=decay,
            )
            score = run_question2(january, config, cold_start).metrics["total_cost_yuan"]
            scores[f"count={count},decay={decay:.2f}"] = score
            candidates.append((score, count, abs(decay - 1.0), config))
    return min(candidates, key=lambda item: item[:3])[3], scores


def run_question2(data: YearData, config: Question2Config, cold_start: ColdStartForecast) -> YearResult:
    if config.planner not in {"deterministic", "stochastic"}:
        raise ValueError("planner must be deterministic or stochastic")
    started = time.perf_counter()
    carried_soc = INITIAL_SOC_KWH
    records: list[DailyRecord] = []
    history_dates: list[date] = []
    history_actual_load: list[np.ndarray] = []
    history_actual_pv: list[np.ndarray] = []
    history_load_residual: list[np.ndarray] = []
    history_pv_residual: list[np.ndarray] = []
    for index, run_date in enumerate(data.dates):
        forecast = forecast_day(
            data,
            index,
            config.forecast,
            cold_start_load_kw=cold_start.load_kw,
            cold_start_pv_kw=cold_start.pv_kw,
        )
        expected_emergency_cost = 0.0
        if config.planner == "stochastic" and history_dates:
            scenarios = build_joint_residual_scenarios(
                history_dates=tuple(history_dates),
                actual_load_kw=np.stack(history_actual_load),
                actual_pv_kw=np.stack(history_actual_pv),
                load_residual_kw=np.stack(history_load_residual),
                pv_residual_kw=np.stack(history_pv_residual),
                target_forecast=forecast,
                decision_date=run_date,
                max_scenarios=config.scenario_count,
                decay=config.scenario_decay,
                window_days=config.scenario_window_days,
            )
            stochastic = plan_two_stage_stochastic(
                run_date,
                scenarios,
                data.price_yuan_per_kwh,
                carried_soc,
                config.reserve_kwh,
                config.emergency_multiplier,
            )
            plan = stochastic.plan
            expected_emergency_cost = stochastic.expected_emergency_cost_yuan
        else:
            plan = plan_day_ahead(run_date, forecast, data.price_yuan_per_kwh, carried_soc, config.reserve_kwh)
        realized = RealizedDay(run_date, data.load_kwh[index].copy(), data.pv_kwh[index].copy())
        execution = simulate_fixed_plan(
            plan,
            realized,
            carried_soc,
            emergency_multiplier=config.emergency_multiplier,
            reserve_kwh=config.reserve_kwh,
        )
        records.append(DailyRecord(
            run_date,
            forecast,
            plan,
            execution,
            data.load_kw[index].copy(),
            data.pv_kw[index].copy(),
            expected_emergency_cost,
        ))
        history_dates.append(run_date)
        history_actual_load.append(data.load_kw[index].copy())
        history_actual_pv.append(data.pv_kw[index].copy())
        history_load_residual.append(data.load_kw[index] - forecast.load_kw)
        history_pv_residual.append(data.pv_kw[index] - forecast.pv_kw)
        carried_soc = float(execution.soc_kwh[-1])
    metrics = evaluate_year(tuple(records))
    return YearResult(
        days=tuple(records),
        metrics=metrics,
        run_metadata={
            "forecast_method": config.forecast.method,
            "planner": config.planner,
            "planning_quantile": config.forecast.planning_quantile,
            "scenario_count": config.scenario_count,
            "scenario_decay": config.scenario_decay,
            "execution_strategy": "fixed_plan",
            "elapsed_seconds": time.perf_counter() - started,
            "simulated_days": len(records),
        },
    )


def evaluate_year(days: tuple[DailyRecord, ...]) -> dict[str, float]:
    if not days:
        raise ValueError("At least one daily record is required")
    official = tuple(day for day in days if day.date >= OFFICIAL_START) or days
    actual_load = np.stack([day.actual_load_kw for day in official])
    actual_pv = np.stack([day.actual_pv_kw for day in official])
    predicted_load = np.stack([day.forecast.load_kw for day in official])
    predicted_pv = np.stack([day.forecast.pv_kw for day in official])
    load_metrics = forecast_metrics(actual_load, predicted_load)
    pv_metrics = forecast_metrics(actual_pv, predicted_pv)
    net_metrics = forecast_metrics(actual_load - actual_pv, predicted_load - predicted_pv)
    emergency = float(sum(np.sum(day.execution.emergency_kwh) for day in official))
    planned_cost = float(sum(day.plan.planned_cost_yuan for day in official))
    emergency_cost = float(sum(day.execution.emergency_cost_yuan for day in official))
    curtailment = float(sum(np.sum(day.execution.curtailment_kwh) for day in official))
    return {
        "load_mae_kw": load_metrics["mae"],
        "load_rmse_kw": load_metrics["rmse"],
        "load_bias_kw": load_metrics["bias"],
        "pv_mae_kw": pv_metrics["mae"],
        "pv_rmse_kw": pv_metrics["rmse"],
        "pv_bias_kw": pv_metrics["bias"],
        "net_mae_kw": net_metrics["mae"],
        "net_rmse_kw": net_metrics["rmse"],
        "net_bias_kw": net_metrics["bias"],
        "planned_purchase_kwh": float(sum(np.sum(day.plan.grid_kwh) for day in official)),
        "planned_cost_yuan": planned_cost,
        "emergency_purchase_kwh": emergency,
        "emergency_cost_yuan": emergency_cost,
        "scenario_expected_emergency_cost_yuan": float(sum(day.expected_emergency_cost_yuan for day in official)),
        "total_cost_yuan": planned_cost + emergency_cost,
        "charge_kwh": float(sum(np.sum(day.execution.charge_kwh) for day in official)),
        "discharge_kwh": float(sum(np.sum(day.execution.discharge_kwh) for day in official)),
        # In fixed-plan settlement this is total surplus energy, not pure PV
        # curtailment, because an over-purchased grid plan can also contribute.
        "surplus_energy_kwh": curtailment,
        "emergency_event_count": float(sum(len(compress_emergency_events(day.execution.emergency_kwh)) for day in official)),
        "soc_min_kwh": float(min(np.min(day.execution.soc_kwh) for day in official)),
        "soc_max_kwh": float(max(np.max(day.execution.soc_kwh) for day in official)),
    }


def _copy_row_style(sheet, source_row: int, target_row: int, max_column: int) -> None:
    sheet.row_dimensions[target_row].height = sheet.row_dimensions[source_row].height
    for column in range(1, max_column + 1):
        source = sheet.cell(source_row, column)
        target = sheet.cell(target_row, column)
        if source.has_style:
            target._style = copy(source._style)
        target.number_format = source.number_format
        target.alignment = copy(source.alignment)
        target.border = copy(source.border)
        target.fill = copy(source.fill)
        target.font = copy(source.font)


def _clock(minute: int) -> str:
    if minute == 1440:
        return "24:00"
    return f"{minute // 60}:{minute % 60:02d}"


def _event_interval(start: int, end: int) -> str:
    return f"{_clock(start)}-{_clock(end)}"


def write_result2_workbook(template: Path, output: Path, result: YearResult) -> None:
    """Fill the official template, validate it, then replace the destination atomically."""

    official = result.official_days
    if len(official) != 334 or official[0].date != OFFICIAL_START or official[-1].date != date(2025, 12, 31):
        raise ValueError("Official output requires every day from February 1 through December 31")
    workbook = load_workbook(template)
    if workbook.sheetnames != ["计划购电量", "充放电量", "紧急购电量"]:
        raise ValueError("Unexpected result2 workbook sheet structure")

    purchase = workbook["计划购电量"]
    for row, record in enumerate(official, start=2):
        purchase.cell(row, 1).value = record.date
        values = internal_to_template(record.plan.grid_kwh)
        for column, value in enumerate(values, start=2):
            purchase.cell(row, column).value = float(value)
        purchase.cell(row, 146).value = float(np.sum(record.plan.grid_kwh))
        purchase.cell(row, 147).value = record.plan.planned_cost_yuan

    storage = workbook["充放电量"]
    style_rows = [[copy(storage.cell(source, column)._style) for column in range(1, 7)] for source in range(2, 8)]
    style_formats = [[storage.cell(source, column).number_format for column in range(1, 7)] for source in range(2, 8)]
    if storage.max_row > 1:
        storage.delete_rows(2, storage.max_row - 1)
    blocks = tuple((start, start + 24) for start in range(0, 144, 24))
    output_row = 2
    for record in official:
        for block_index, (start, end) in enumerate(blocks):
            for column in range(1, 7):
                storage.cell(output_row, column)._style = copy(style_rows[block_index][column - 1])
                storage.cell(output_row, column).number_format = style_formats[block_index][column - 1]
            storage.cell(output_row, 1).value = record.date if block_index == 0 else None
            storage.cell(output_row, 2).value = f"{start // 6}:00-{end // 6}:00"
            storage.cell(output_row, 3).value = float(np.sum(record.execution.charge_kwh[start:end]))
            storage.cell(output_row, 4).value = float(np.sum(record.execution.discharge_kwh[start:end]))
            if block_index == 0:
                storage.cell(output_row, 5).value = "0:00"
                storage.cell(output_row, 6).value = float(record.execution.soc_kwh[0])
            elif block_index == 1:
                storage.cell(output_row, 5).value = "24:00"
                storage.cell(output_row, 6).value = float(record.execution.soc_kwh[-1])
            output_row += 1

    emergency = workbook["紧急购电量"]
    emergency_styles = [[copy(emergency.cell(source, column)._style) for column in range(1, 4)] for source in range(2, 5)]
    emergency_formats = [[emergency.cell(source, column).number_format for column in range(1, 4)] for source in range(2, 5)]
    if emergency.max_row > 1:
        emergency.delete_rows(2, emergency.max_row - 1)
    output_row = 2
    for record in official:
        events = compress_emergency_events(record.execution.emergency_kwh)
        rows = events or (None,)
        for event_index, event in enumerate(rows):
            style_index = min(event_index, 2)
            for column in range(1, 4):
                emergency.cell(output_row, column)._style = copy(emergency_styles[style_index][column - 1])
                emergency.cell(output_row, column).number_format = emergency_formats[style_index][column - 1]
            emergency.cell(output_row, 1).value = record.date if event_index == 0 else None
            if event is not None:
                emergency.cell(output_row, 2).value = _event_interval(event.start_minute, event.end_minute)
                emergency.cell(output_row, 3).value = event.energy_kwh
            output_row += 1

    output.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary_name = tempfile.mkstemp(suffix=".xlsx", dir=output.parent)
    os.close(descriptor)
    temporary = Path(temporary_name)
    try:
        workbook.save(temporary)
        workbook.close()
        _validate_result2(temporary, result)
        os.replace(temporary, output)
    finally:
        temporary.unlink(missing_ok=True)


def _validate_result2(path: Path, result: YearResult) -> None:
    workbook = load_workbook(path, data_only=True, read_only=True)
    if workbook.sheetnames != ["计划购电量", "充放电量", "紧急购电量"]:
        raise RuntimeError("Result2 sheet names changed")
    purchase = workbook["计划购电量"]
    if purchase.max_row != 335 or purchase.max_column != 147:
        raise RuntimeError("Result2 purchase sheet has unexpected dimensions")
    rows = purchase.iter_rows(min_row=2, values_only=True)
    for row, (values_row, record) in enumerate(zip(rows, result.official_days, strict=True), start=2):
        workbook_date = values_row[0].date() if hasattr(values_row[0], "date") else values_row[0]
        if workbook_date != record.date:
            raise RuntimeError(f"Result2 date mismatch on row {row}")
        values = np.asarray(values_row[1:145], dtype=float)
        if not np.allclose(values, internal_to_template(record.plan.grid_kwh), atol=1e-6):
            raise RuntimeError(f"Result2 purchase mismatch on row {row}")
        if abs(float(values_row[145]) - float(np.sum(record.plan.grid_kwh))) > 1e-5:
            raise RuntimeError(f"Result2 daily total mismatch on row {row}")
        if abs(float(values_row[146]) - record.plan.planned_cost_yuan) > 1e-5:
            raise RuntimeError(f"Result2 daily cost mismatch on row {row}")

    storage = workbook["充放电量"]
    if storage.max_row != 1 + 6 * len(result.official_days):
        raise RuntimeError("Result2 storage sheet does not contain six rows per day")
    storage_rows = storage.iter_rows(min_row=2, values_only=True)
    for record in result.official_days:
        group = [next(storage_rows) for _ in range(6)]
        workbook_date = group[0][0].date() if hasattr(group[0][0], "date") else group[0][0]
        if workbook_date != record.date:
            raise RuntimeError(f"Result2 storage date mismatch for {record.date}")
        if abs(sum(float(row[2]) for row in group) - float(np.sum(record.execution.charge_kwh))) > 1e-5:
            raise RuntimeError(f"Result2 charge total mismatch for {record.date}")
        if abs(sum(float(row[3]) for row in group) - float(np.sum(record.execution.discharge_kwh))) > 1e-5:
            raise RuntimeError(f"Result2 discharge total mismatch for {record.date}")
        if abs(float(group[0][5]) - float(record.execution.soc_kwh[0])) > 1e-5:
            raise RuntimeError(f"Result2 initial SOC mismatch for {record.date}")
        if abs(float(group[1][5]) - float(record.execution.soc_kwh[-1])) > 1e-5:
            raise RuntimeError(f"Result2 terminal SOC mismatch for {record.date}")

    emergency = workbook["紧急购电量"]
    written_emergency = sum(
        float(row[2]) for row in emergency.iter_rows(min_row=2, values_only=True) if row[2] is not None
    )
    expected_emergency = sum(float(np.sum(record.execution.emergency_kwh)) for record in result.official_days)
    if abs(written_emergency - expected_emergency) > 1e-5:
        raise RuntimeError("Result2 emergency-purchase total mismatch")
    workbook.close()


def write_summary_workbook(output: Path, result: YearResult, comparisons: dict[str, YearResult] | None = None) -> None:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "策略汇总"
    summary.append(["指标", "数值"])
    for key, value in result.metrics.items():
        summary.append([key, value])
    comparison = workbook.create_sheet("方案对比")
    comparison.append(["方案", "净负荷RMSE/kW", "计划费用/元", "紧急购电费/元", "总费用/元", "弃电量/kWh"])
    for name, case in (comparisons or {"official": result}).items():
        comparison.append([
            name,
            case.metrics["net_rmse_kw"],
            case.metrics["planned_cost_yuan"],
            case.metrics["emergency_cost_yuan"],
            case.metrics["total_cost_yuan"],
            case.metrics["surplus_energy_kwh"],
        ])
    specified = workbook.create_sheet("指定日期")
    specified.append(["日期", "计划购电量/kWh", "计划费用/元", "紧急购电量/kWh", "紧急购电费/元", "总费用/元"])
    for record in result.official_days:
        if record.date in SPECIFIED_DATES:
            specified.append([
                record.date,
                float(np.sum(record.plan.grid_kwh)),
                record.plan.planned_cost_yuan,
                float(np.sum(record.execution.emergency_kwh)),
                record.execution.emergency_cost_yuan,
                record.execution.total_cost_yuan,
            ])
    for sheet in workbook.worksheets:
        sheet.sheet_view.showGridLines = False
        sheet.freeze_panes = "A2"
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.font = Font(name="Arial", size=10)
                cell.alignment = Alignment(vertical="center")
                if isinstance(cell.value, float):
                    cell.number_format = "0.0000"
        for column in range(1, sheet.max_column + 1):
            values = [str(sheet.cell(row, column).value or "") for row in range(1, min(sheet.max_row, 200) + 1)]
            sheet.column_dimensions[get_column_letter(column)].width = min(max(max(map(len, values)) + 2, 12), 28)
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)


def plot_question2(output_directory: Path, result: YearResult) -> tuple[Path, Path]:
    os.environ.setdefault("MPLCONFIGDIR", str(Path(tempfile.gettempdir()) / "cumcm-matplotlib"))
    import matplotlib

    matplotlib.use("Agg", force=True)
    import matplotlib.pyplot as plt

    plt.rcParams["font.sans-serif"] = ["Microsoft YaHei", "SimHei", "DejaVu Sans"]
    plt.rcParams["axes.unicode_minus"] = False
    output_directory.mkdir(parents=True, exist_ok=True)
    specified_path = output_directory / "q2_specified_days.pdf"
    monthly_path = output_directory / "q2_monthly_cost.pdf"

    by_date = {record.date: record for record in result.official_days}
    hours = np.arange(144) / 6.0
    figure, axes = plt.subplots(2, 2, figsize=(10.8, 6.8), sharex=True)
    for axis, run_date in zip(axes.flat, SPECIFIED_DATES):
        record = by_date[run_date]
        actual_net = record.actual_load_kw - record.actual_pv_kw
        predicted_net = record.forecast.load_kw - record.forecast.pv_kw
        axis.plot(hours, actual_net, color="#1f4e79", linewidth=1.2, label="实际净负荷")
        axis.plot(hours, predicted_net, color="#70ad47", linewidth=1.0, linestyle="--", label="点预测")
        axis.plot(hours, record.forecast.planning_net_kw, color="#c55a11", linewidth=1.0, label="80%分位计划曲线")
        axis.set_title(run_date.strftime("%Y-%m-%d"))
        axis.set_xlim(0, 24)
        axis.grid(alpha=0.25)
    axes[0, 0].set_ylabel("功率/kW")
    axes[1, 0].set_ylabel("功率/kW")
    axes[1, 0].set_xlabel("时刻/h")
    axes[1, 1].set_xlabel("时刻/h")
    handles, labels = axes[0, 0].get_legend_handles_labels()
    figure.legend(handles, labels, loc="upper center", ncol=3, frameon=False)
    figure.tight_layout(rect=(0, 0, 1, 0.94))
    figure.savefig(specified_path, bbox_inches="tight")
    plt.close(figure)

    months = tuple(range(2, 13))
    normal = []
    emergency = []
    for month in months:
        selected = [record for record in result.official_days if record.date.month == month]
        normal.append(sum(record.plan.planned_cost_yuan for record in selected) / 10_000.0)
        emergency.append(sum(record.execution.emergency_cost_yuan for record in selected) / 10_000.0)
    figure, axis = plt.subplots(figsize=(9.2, 4.4))
    axis.bar(months, normal, color="#4472c4", label="计划购电费")
    axis.bar(months, emergency, bottom=normal, color="#ed7d31", label="紧急购电费")
    axis.set_xticks(months)
    axis.set_xlabel("月份")
    axis.set_ylabel("费用/万元")
    axis.grid(axis="y", alpha=0.25)
    axis.legend(frameon=False, ncol=2)
    figure.tight_layout()
    figure.savefig(monthly_path, bbox_inches="tight")
    plt.close(figure)
    return specified_path, monthly_path


def _project_root() -> Path:
    return Path(__file__).resolve().parents[2]


def main() -> None:
    root = _project_root()
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--attachment1", type=Path, default=root / "problems/C题/附件/附件1.xlsx")
    parser.add_argument("--attachment2", type=Path, default=root / "problems/C题/附件/附件2.xlsx")
    parser.add_argument("--template", type=Path, default=root / "src/附件5/result2.xlsx")
    parser.add_argument("--output", type=Path, default=root / "src/附件5/result2.xlsx")
    parser.add_argument("--summary", type=Path, default=root / "src/data/question2_summary.xlsx")
    parser.add_argument("--forecast", choices=("previous_day", "seven_day", "week_type", "similar_day"), default="seven_day")
    parser.add_argument("--decay", type=float, default=0.95)
    parser.add_argument("--candidate-count", type=int, default=14)
    parser.add_argument("--quantile", type=float, default=0.80)
    parser.add_argument("--planner", choices=("deterministic", "stochastic"), default="deterministic")
    parser.add_argument("--scenario-count", type=int, default=14)
    parser.add_argument("--scenario-decay", type=float, default=0.95)
    parser.add_argument("--scenario-window", type=int, default=90)
    parser.add_argument("--calibrate-scenarios", action="store_true")
    parser.add_argument("--write-results", action="store_true")
    parser.add_argument("--plots", action="store_true")
    parser.add_argument("--compare", action="store_true")
    args = parser.parse_args()
    data = load_year_data(args.attachment1, args.attachment2)
    cold_start = load_cold_start_forecast(args.attachment1)
    config = Question2Config(
        forecast=ForecastConfig(
            method=args.forecast,
            candidate_count=args.candidate_count,
            lambda_load=args.decay,
            lambda_pv=args.decay,
            planning_quantile=args.quantile,
        ),
        planner=args.planner,
        scenario_count=args.scenario_count,
        scenario_decay=args.scenario_decay,
        scenario_window_days=args.scenario_window,
    )
    if args.calibrate_scenarios:
        if config.planner != "stochastic":
            parser.error("--calibrate-scenarios requires --planner stochastic")
        config, calibration_scores = calibrate_stochastic_parameters(data, cold_start, config)
        for name, score in calibration_scores.items():
            print(f"january_calibration[{name}]={score:.6f}")
        print(f"selected_scenario_count={config.scenario_count}")
        print(f"selected_scenario_decay={config.scenario_decay:.2f}")
    result = run_question2(data, config, cold_start)
    print(f"simulated_days={len(result.days)} official_days={len(result.official_days)}")
    for key, value in result.metrics.items():
        print(f"{key}={value:.6f}")
    print(f"elapsed_seconds={result.run_metadata['elapsed_seconds']:.3f}")
    if args.write_results:
        write_result2_workbook(args.template, args.output, result)
        comparisons: dict[str, YearResult] | None = None
        if args.compare:
            comparisons = {f"{args.forecast}_{config.planner}": result}
            for method in ("previous_day", "seven_day", "week_type", "similar_day"):
                comparison_name = f"{method}_deterministic"
                if comparison_name not in comparisons:
                    comparisons[comparison_name] = run_question2(
                        data,
                        Question2Config(forecast=ForecastConfig(method=method)),
                        cold_start,
                    )
        write_summary_workbook(args.summary, result, comparisons)
        print(f"result2={args.output}")
        print(f"summary={args.summary}")
    if args.plots:
        paths = plot_question2(root / "src/tex/figure", result)
        print("plots=" + ",".join(str(path) for path in paths))


if __name__ == "__main__":
    main()
